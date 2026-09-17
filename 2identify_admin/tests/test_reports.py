"""XLSX exports from the same authenticated alert API used by the Admin."""

from __future__ import annotations

from datetime import UTC, date, datetime
from io import BytesIO
from time import monotonic, sleep
from uuid import uuid4

from openpyxl import load_workbook
from PySide6.QtTest import QSignalSpy

from app.controllers.reports_controller import ReportsController
from app.core.session import AdminSessionContext
from app.domain import AdminAuthentication, Administrator
from app.domain.alerts import (
    AdminAlert,
    AdminAlertPage,
    AlertCamera,
    AlertEmployee,
    AlertOccurrence,
    AlertOperationalContext,
    AlertSector,
)
from app.services.admin_reports_service import AdminReportsService, ReportFilters
from app.ui.main.main_window import MainWindow
from app.ui.reports import ReportsPage


def _alert(alert_id: int, camera_id: int, *, category: str = "ppe") -> AdminAlert:
    now = datetime(2026, 9, 17, 10, 30, tzinfo=UTC)
    sector = AlertSector(2, "Usinagem")
    return AdminAlert(
        id=alert_id,
        category=category,
        level="critical",
        status="nao_lido",
        summary="Capacete na mão" if category == "ppe" else "Pessoa em área de risco",
        observation="=Atenção à peça",
        created_at=now,
        received_at=now,
        confirmed_at=None,
        confirmed_by=None,
        closed_at=None,
        closed_by=None,
        occurrence=AlertOccurrence(
            id=alert_id + 100,
            type="ppe_absent" if category == "ppe" else "person_in_risk_area",
            description="Ocorrência de João",
            confidence=0.91,
            image_reference=None,
            video_reference=None,
            detected_at=now,
            employee=AlertEmployee(7, "João Ávila", "MAT-7", "Operador", "Manhã", sector),
            camera=AlertCamera(camera_id, f"Fresa {camera_id}", None, sector),
        ),
        operational_context=AlertOperationalContext(
            uuid4(), uuid4(), 12, None, "ppe_absent", "ppe:1", None, now
        ),
    )


class _Provider:
    def __init__(self, alerts: tuple[AdminAlert, ...]) -> None:
        self.alerts = alerts
        self.offsets: list[int] = []

    def get_alerts(self, token: str, *, limit: int, offset: int) -> AdminAlertPage:
        assert token == "session-token"
        self.offsets.append(offset)
        return AdminAlertPage(self.alerts[offset: offset + limit], len(self.alerts), limit, offset)


def test_export_paginates_and_keeps_camera_identity_and_dates() -> None:
    provider = _Provider((_alert(1, 5), _alert(2, 8), _alert(3, 5, category="risk_area")))
    content = AdminReportsService(provider, page_size=2).export_xlsx(
        "session-token", ReportFilters(date_from=date(2026, 9, 17))
    )
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["Resumo", "Alertas", "Ocorrências", "EPIs"]
    assert provider.offsets == [0, 2]
    alerts = workbook["Alertas"]
    assert alerts["A1"].value == "ID alerta"
    assert [alerts.cell(row, 10).value for row in (2, 3, 4)] == [5, 8, 5]
    assert alerts["C2"].value == datetime(2026, 9, 17, 10, 30)
    assert alerts.auto_filter.ref == "A1:Q4"
    assert alerts.freeze_panes == "A2"
    assert workbook["EPIs"].max_row == 3
    assert workbook["Ocorrências"]["D2"].value == "João Ávila"
    assert workbook["Resumo"]["B4"].value == 3
    assert alerts["H2"].value == "'=Atenção à peça"


def test_report_filters_and_empty_workbook() -> None:
    provider = _Provider((_alert(1, 5), _alert(2, 8)))
    content = AdminReportsService(provider).export_xlsx(
        "session-token", ReportFilters(camera_id=8, severity="critical")
    )
    workbook = load_workbook(BytesIO(content))
    assert workbook["Alertas"].max_row == 2
    assert workbook["Alertas"].cell(2, 10).value == 8
    assert workbook["Resumo"]["B4"].value == 1
    empty = AdminReportsService(_Provider(())).export_xlsx("session-token", ReportFilters())
    workbook = load_workbook(BytesIO(empty))
    assert workbook["Alertas"].max_row == 1
    assert workbook["Ocorrências"].max_row == 1
    assert workbook["EPIs"].max_row == 1


def test_report_rejects_inverted_dates() -> None:
    try:
        ReportFilters(date_from=date(2026, 9, 18), date_to=date(2026, 9, 17))
    except ValueError:
        pass
    else:
        raise AssertionError("intervalo invertido precisa ser rejeitado")


def test_reports_page_is_real_route_and_emits_selected_filters(qapp) -> None:
    admin = Administrator(1, "Admin", "admin", "administrador")
    window = MainWindow(admin)
    window.trocar_pagina("relatorios")
    assert window.stack.currentWidget() is window.reports
    spy = QSignalSpy(window.reports.export_requested)
    window.reports.camera_id.setValue(8)
    window.reports.export_button.click()
    assert spy.count() == 1
    assert spy.at(0)[0].camera_id == 8
    window.close()


def test_report_controller_saves_valid_xlsx_without_blocking_ui(
    qapp, tmp_path, monkeypatch
) -> None:
    page = ReportsPage()
    target = tmp_path / "relatorio.xlsx"
    monkeypatch.setattr(
        "app.controllers.reports_controller.QFileDialog.getSaveFileName",
        lambda *args: (str(target), "Planilha Excel (*.xlsx)"),
    )
    session = AdminSessionContext()
    session.open(
        AdminAuthentication(
            Administrator(1, "Admin", "admin", "administrador"),
            "token", expires_in=60,
        )
    )
    provider = _Provider((_alert(1, 5),))
    # The provider must see the exact in-memory session token, never a database connection.
    provider.get_alerts = lambda token, *, limit, offset: AdminAlertPage(
        (_alert(1, 5),) if offset == 0 and token == "token" else (), 1, limit, offset
    )
    controller = ReportsController(
        page, AdminReportsService(provider), session, shutdown_timeout_ms=2000
    )
    page.export_button.click()
    deadline = monotonic() + 3
    while not target.is_file() and monotonic() < deadline:
        qapp.processEvents()
        sleep(0.01)
    assert target.is_file()
    assert load_workbook(target)["Alertas"].max_row == 2
    assert controller.shutdown()
    page.close()
