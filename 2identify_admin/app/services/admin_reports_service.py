"""Generate a presentable XLSX from paginated, authenticated Admin API data."""

from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, date, datetime
from io import BytesIO
from typing import Literal, Protocol

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.alerts import AdminAlert, AdminAlertPage
from app.services.errors import InvalidApiResponseError


@dataclass(frozen=True, slots=True)
class ReportFilters:
    date_from: date | None = None
    date_to: date | None = None
    sector_id: int | None = None
    employee_id: int | None = None
    operation_id: int | None = None
    camera_id: int | None = None
    severity: Literal["critical", "warning"] | None = None

    def __post_init__(self) -> None:
        if self.date_from and self.date_to and self.date_from > self.date_to:
            raise ValueError("A data inicial deve ser anterior à data final.")
        for value in (self.sector_id, self.employee_id, self.operation_id, self.camera_id):
            if value is not None and value <= 0:
                raise ValueError("IDs de filtros devem ser positivos.")


class ReportsAlertProvider(Protocol):
    def get_alerts(
        self, access_token: str, *, limit: int, offset: int
    ) -> AdminAlertPage: ...


class AdminReportsService:
    """Read through the existing API and export only fields actually available."""

    def __init__(self, provider: ReportsAlertProvider, *, page_size: int = 100) -> None:
        if not 1 <= page_size <= 100:
            raise ValueError("page_size deve estar entre 1 e 100")
        self._provider = provider
        self._page_size = page_size

    def export_xlsx(self, access_token: str, filters: ReportFilters) -> bytes:
        alerts = tuple(item for item in self._all_alerts(access_token) if _matches(item, filters))
        workbook = Workbook()
        summary = workbook.active
        assert summary is not None
        summary.title = "Resumo"
        _summary(summary, alerts, filters)
        _table(
            workbook.create_sheet("Alertas"),
            (
                "ID alerta", "ID ocorrência", "Criado em (UTC)", "Severidade", "Status",
                "Categoria", "Resumo", "Observação", "Funcionário", "ID câmera",
                "Câmera", "Setor", "ID operação", "Tipo ocorrência",
                "Detectado em (UTC)", "EPI", "Conformidade",
            ),
            (
                (
                    alert.id, alert.occurrence.id, _excel_date(alert.created_at),
                    alert.level, alert.status, alert.category, alert.summary,
                    alert.observation, _employee_name(alert), _camera_id(alert),
                    _camera_name(alert), _sector_name(alert), _operation_id(alert),
                    alert.occurrence.type, _excel_date(alert.occurrence.detected_at),
                    _ppe_subject(alert), _conformity(alert),
                ) for alert in alerts
            ),
            date_columns=(3, 15),
        )
        _table(
            workbook.create_sheet("Ocorrências"),
            (
                "ID ocorrência", "Tipo", "Descrição", "Funcionário", "ID câmera",
                "Câmera", "Setor", "Detectado em (UTC)", "Confiança", "ID alerta",
            ),
            (
                (
                    alert.occurrence.id, alert.occurrence.type,
                    alert.occurrence.description, _employee_name(alert), _camera_id(alert),
                    _camera_name(alert), _sector_name(alert),
                    _excel_date(alert.occurrence.detected_at), alert.occurrence.confidence,
                    alert.id,
                ) for alert in alerts
            ),
            date_columns=(8,),
        )
        _table(
            workbook.create_sheet("EPIs"),
            (
                "ID alerta", "EPI", "ID câmera", "Câmera", "Funcionário",
                "Observado em (UTC)", "Conformidade", "Status",
            ),
            (
                (
                    alert.id, _ppe_subject(alert), _camera_id(alert),
                    _camera_name(alert), _employee_name(alert),
                    _excel_date(alert.occurrence.detected_at), "Não conforme", alert.status,
                ) for alert in alerts if alert.category == "ppe"
            ),
            date_columns=(6,),
        )
        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def _all_alerts(self, access_token: str) -> tuple[AdminAlert, ...]:
        collected: list[AdminAlert] = []
        offset = 0
        while True:
            page = self._provider.get_alerts(access_token, limit=self._page_size, offset=offset)
            if page.offset != offset or page.limit != self._page_size:
                raise InvalidApiResponseError("Paginação de alertas inconsistente.")
            if not page.items:
                if offset < page.total:
                    raise InvalidApiResponseError("A API interrompeu a paginação dos alertas.")
                break
            collected.extend(page.items)
            offset += len(page.items)
            if offset >= page.total:
                break
        return tuple(collected)


def _matches(alert: AdminAlert, filters: ReportFilters) -> bool:
    created = alert.created_at.date()
    employee = alert.occurrence.employee
    camera = alert.occurrence.camera
    context = alert.operational_context
    sector = employee.sector if employee and employee.sector else camera.sector if camera else None
    return (
        (filters.date_from is None or created >= filters.date_from)
        and (filters.date_to is None or created <= filters.date_to)
        and (filters.sector_id is None or (sector is not None and sector.id == filters.sector_id))
        and (
            filters.employee_id is None
            or (employee is not None and employee.id == filters.employee_id)
        )
        and (
            filters.operation_id is None
            or (context is not None and context.operation_id == filters.operation_id)
        )
        and (filters.camera_id is None or (camera is not None and camera.id == filters.camera_id))
        and (filters.severity is None or alert.level == filters.severity)
    )


def _excel_date(value: datetime) -> datetime:
    return value.astimezone(UTC).replace(tzinfo=None) if value.tzinfo else value


def _employee_name(alert: AdminAlert) -> str | None:
    employee = alert.occurrence.employee
    return employee.name if employee else None


def _camera_id(alert: AdminAlert) -> int | None:
    camera = alert.occurrence.camera
    return camera.id if camera else None


def _camera_name(alert: AdminAlert) -> str | None:
    camera = alert.occurrence.camera
    return camera.name if camera else None


def _sector_name(alert: AdminAlert) -> str | None:
    employee = alert.occurrence.employee
    camera = alert.occurrence.camera
    sector = employee.sector if employee and employee.sector else camera.sector if camera else None
    return sector.name if sector else None


def _operation_id(alert: AdminAlert) -> int | None:
    context = alert.operational_context
    return context.operation_id if context else None


def _ppe_subject(alert: AdminAlert) -> str | None:
    if alert.category != "ppe":
        return None
    context = alert.operational_context
    return context.subject_key if context else alert.occurrence.description


def _conformity(alert: AdminAlert) -> str | None:
    return "Não conforme" if alert.category == "ppe" else None


def _safe(value: object) -> object:
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _table(
    sheet: Worksheet,
    headers: tuple[str, ...],
    rows: Iterable[tuple[object, ...]],
    *,
    date_columns: tuple[int, ...] = (),
) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(tuple(_safe(value) for value in row))
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for column in range(1, len(headers) + 1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = min(40, max(15, len(headers[column - 1]) + 3))
    for column in date_columns:
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row, column).number_format = "dd/mm/yyyy hh:mm:ss"
        sheet.column_dimensions[get_column_letter(column)].width = 21
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"


def _summary(sheet: Worksheet, alerts: tuple[AdminAlert, ...], filters: ReportFilters) -> None:
    sheet.append(("Relatório 2Identify", "Resumo de alertas registrados"))
    sheet.append(("Gerado em (UTC)", datetime.now(UTC).replace(tzinfo=None)))
    sheet.append(("Período", f"{filters.date_from or 'Início'} a {filters.date_to or 'Hoje'}"))
    sheet.append(("Total de alertas", len(alerts)))
    sheet.append(("Críticos", sum(item.level == "critical" for item in alerts)))
    sheet.append(("Alertas de EPI", sum(item.category == "ppe" for item in alerts)))
    camera_ids = {_camera_id(item) for item in alerts if _camera_id(item) is not None}
    sheet.append(("Câmeras distintas", len(camera_ids)))
    sheet["A1"].font = Font(bold=True, color="17365D", size=16)
    sheet["B2"].number_format = "dd/mm/yyyy hh:mm:ss"
    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 48
